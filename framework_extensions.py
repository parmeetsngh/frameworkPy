"""
Extensions for the Test Automation Framework that provide:
1. Better document parsing (using libraries)
2. NLP-based test case generation
3. Traceability reporting
"""

import re
import json
import logging
from typing import List, Dict, Optional, Set
from pathlib import Path
import pandas as pd
import requests

# Note: These imports would need to be installed in a real implementation
# pip install python-docx pandas openpyxl nltk spacy gensim

try:
    import docx
    import pandas as pd
    import nltk
    import spacy
    from gensim.summarization import keywords

    # Download required NLTK data
    nltk.download('punkt', quiet=True)
    nltk.download('stopwords', quiet=True)
    nltk.download('wordnet', quiet=True)

    # Load spaCy model
    nlp = spacy.load('en_core_web_sm')

    LIBRARIES_LOADED = True
except ImportError:
    LIBRARIES_LOADED = False
    logging.warning("Some libraries couldn't be imported. Advanced features will be disabled.")


class EnhancedDocxParser:
    """Enhanced parser for Word documents using python-docx"""

    def parse(self, file_path: str) -> Dict:
        if not LIBRARIES_LOADED:
            logging.warning("python-docx not available. Using basic parsing.")
            return {"content": f"Content from {file_path}", "sections": []}

        try:
            doc = docx.Document(file_path)

            # Extract document structure
            result = {
                "content": "",
                "sections": [],
                "headings": [],
                "tables": []
            }

            # Process paragraphs
            current_section = None
            for para in doc.paragraphs:
                # Add paragraph text to content
                result["content"] += para.text + "\n"

                # Check if this is a heading
                if para.style.name.startswith('Heading'):
                    heading_level = int(para.style.name.replace('Heading ', ''))
                    result["headings"].append({
                        "level": heading_level,
                        "text": para.text
                    })

                    # Create a new section
                    current_section = {
                        "heading": para.text,
                        "level": heading_level,
                        "content": ""
                    }
                    result["sections"].append(current_section)
                elif current_section:
                    current_section["content"] += para.text + "\n"

            # Process tables
            for table in doc.tables:
                table_data = []
                for row in table.rows:
                    row_data = [cell.text for cell in row.cells]
                    table_data.append(row_data)

                result["tables"].append(table_data)

            return result

        except Exception as e:
            logging.error(f"Error parsing DOCX file {file_path}: {e}")
            return {"error": str(e)}


class EnhancedExcelParser:
    """Enhanced parser for Excel files using pandas"""

    def parse(self, file_path: str) -> Dict:
        if not LIBRARIES_LOADED:
            logging.warning("pandas not available. Using basic parsing.")
            return {"sheets": [{"name": "Sheet1", "data": [["Dummy data"]]}]}

        try:
            # Read all sheets
            excel_file = pd.ExcelFile(file_path)
            sheet_names = excel_file.sheet_names

            result = {
                "file_path": file_path,
                "sheets": []
            }

            for sheet_name in sheet_names:
                df = pd.read_excel(excel_file, sheet_name)

                # Convert dataframe to dict for easier processing
                sheet_data = {
                    "name": sheet_name,
                    "headers": list(df.columns),
                    "data": df.to_dict(orient='records')
                }

                # Extract test cases if this looks like a test case sheet
                if any(col.lower() in ["test case", "test_case", "testcase", "tc_id", "test id"]
                       for col in df.columns):
                    sheet_data["test_cases"] = self._extract_test_cases(df)

                result["sheets"].append(sheet_data)

            return result

        except Exception as e:
            logging.error(f"Error parsing Excel file {file_path}: {e}")
            return {"error": str(e)}

    def _extract_test_cases(self, df: pd.DataFrame) -> List[Dict]:
        """Extract test cases from a dataframe"""
        test_cases = []

        # Try to identify key columns
        tc_id_col = self._find_column(df, ["test case id", "tc_id", "test id", "testcase id"])
        tc_name_col = self._find_column(df, ["test case name", "tc_name", "test name", "name", "summary"])
        tc_desc_col = self._find_column(df, ["description", "test description", "tc_desc"])
        tc_steps_col = self._find_column(df, ["steps", "test steps", "tc_steps"])
        tc_expected_col = self._find_column(df, ["expected results", "expected", "expectations"])

        # Process each row
        for _, row in df.iterrows():
            test_case = {}

            if tc_id_col:
                test_case["id"] = str(row[tc_id_col])

            if tc_name_col:
                test_case["name"] = str(row[tc_name_col])
            elif "id" in test_case:
                test_case["name"] = f"Test Case {test_case['id']}"
            else:
                test_case["name"] = "Unnamed Test Case"

            if tc_desc_col:
                test_case["description"] = str(row[tc_desc_col])

            # Extract steps
            if tc_steps_col:
                steps_text = str(row[tc_steps_col])
                # Try to split steps if they're in a numbered format
                if re.search(r'^\d+\.', steps_text):
                    steps = re.split(r'\d+\.', steps_text)
                    # Remove empty entries and strip whitespace
                    test_case["steps"] = [step.strip() for step in steps if step.strip()]
                else:
                    # Split by newlines if no numbering is found
                    steps = steps_text.split('\n')
                    test_case["steps"] = [step.strip() for step in steps if step.strip()]

            # Extract expected results
            if tc_expected_col:
                expected_text = str(row[tc_expected_col])
                if re.search(r'^\d+\.', expected_text):
                    expected = re.split(r'\d+\.', expected_text)
                    test_case["expected_results"] = [result.strip() for result in expected if result.strip()]
                else:
                    expected = expected_text.split('\n')
                    test_case["expected_results"] = [result.strip() for result in expected if result.strip()]

            test_cases.append(test_case)

        return test_cases

    def _find_column(self, df: pd.DataFrame, possible_names: List[str]) -> Optional[str]:
        """Find a column in the dataframe based on possible names"""
        for col in df.columns:
            if col.lower() in possible_names:
                return col
        return None


class NLPTestCaseGenerator:
    """Generates test cases using NLP techniques"""

    def __init__(self):
        self.nlp_available = LIBRARIES_LOADED

    def generate_test_scenarios_from_text(self, text: str) -> List[Dict]:
        """Generate test scenarios from raw text using NLP"""
        if not self.nlp_available:
            logging.warning("NLP libraries not available. Using basic test generation.")
            return [self._create_basic_scenario("Basic Test Scenario")]

        try:
            # Process text with spaCy
            doc = nlp(text)

            # Extract sentences
            sentences = [sent.text for sent in doc.sents]

            # Look for requirements and functionalities
            scenarios = []

            # Look for sentences that sound like requirements
            requirement_patterns = [
                r'(must|should|shall|will|needs to)',
                r'(is required to|are required to)',
                r'(functionality|feature|ability)'
            ]

            for sentence in sentences:
                # Check if sentence matches requirement patterns
                if any(re.search(pattern, sentence, re.IGNORECASE) for pattern in requirement_patterns):
                    # Create a scenario from this requirement
                    scenario = self._create_scenario_from_requirement(sentence)
                    scenarios.append(scenario)

            # If no specific requirements found, create some basic scenarios
            if not scenarios:
                # Extract key topics from text
                key_phrases = self._extract_key_phrases(text)
                for phrase in key_phrases[:3]:  # Use top 3 key phrases
                    scenario = self._create_basic_scenario(f"Test {phrase}")
                    scenarios.append(scenario)

            return scenarios

        except Exception as e:
            logging.error(f"Error in NLP test generation: {e}")
            return [self._create_basic_scenario("Basic Test Scenario")]

    def _create_scenario_from_requirement(self, requirement: str) -> Dict:
        """Create a test scenario from a requirement sentence"""
        # Extract the main action from the requirement
        doc = nlp(requirement)

        # Find main verb and its object
        main_verb = None
        main_object = None

        for token in doc:
            if token.pos_ == "VERB" and not main_verb:
                main_verb = token.lemma_
            if token.dep_ in ("dobj", "pobj") and token.pos_ == "NOUN" and not main_object:
                main_object = token.text

        # Create scenario name
        if main_verb and main_object:
            scenario_name = f"Verify {main_verb} {main_object}"
        else:
            # Fallback to using the requirement directly
            scenario_name = f"Verify requirement: {requirement[:50]}..."

        # Create basic steps
        steps = [
            {"type": "Given", "description": "the system is ready"},
            {"type": "When",
             "description": f"the user performs actions related to '{main_object or 'the requirement'}'"},
            {"type": "Then", "description": f"the system should behave according to '{requirement[:50]}...'"}
        ]

        return {
            "name": scenario_name,
            "requirement": requirement,
            "steps": steps,
            "tags": ["automated", "requirement-based"]
        }

    def _create_basic_scenario(self, name: str) -> Dict:
        """Create a basic test scenario with standard steps"""
        return {
            "name": name,
            "steps": [
                {"type": "Given", "description": "the application is running"},
                {"type": "When", "description": "the user performs the required action"},
                {"type": "Then", "description": "the system responds correctly"}
            ],
            "tags": ["automated"]
        }

    def _extract_key_phrases(self, text: str) -> List[str]:
        """Extract key phrases from text using NLP"""
        try:
            # Use gensim's keywords extraction
            key_phrases = keywords(text, words=5, split=True)
            return key_phrases
        except Exception:
            # Fallback to basic noun extraction with spaCy
            doc = nlp(text)
            nouns = [chunk.text for chunk in doc.noun_chunks]
            return nouns[:5]  # Return up to 5 noun phrases


class TraceabilityReporter:
    """Generates traceability reports between requirements and test cases"""

    def __init__(self, output_dir: str):
        self.output_dir = Path(output_dir)

    def generate_traceability_matrix(self, test_scenarios: List[Dict], source_data: Dict) -> Dict:
        """Generate a traceability matrix between requirements and test scenarios"""
        # Extract requirements from source data
        requirements = self._extract_requirements(source_data)

        # Create the matrix
        matrix = {
            "requirements": requirements,
            "test_scenarios": [scenario["name"] for scenario in test_scenarios],
            "coverage": {}
        }

        # Map test scenarios to requirements
        for scenario in test_scenarios:
            scenario_name = scenario["name"]
            matrix["coverage"][scenario_name] = []

            # Check which requirements this scenario covers
            for req_id, req_text in requirements.items():
                # Simple check - see if requirement text appears in scenario or its steps
                scenario_text = json.dumps(scenario).lower()

                # Check if requirement keywords are in the scenario
                req_keywords = self._extract_requirement_keywords(req_text)
                if any(keyword.lower() in scenario_text for keyword in req_keywords):
                    matrix["coverage"][scenario_name].append(req_id)

        # Write the matrix to a file
        self._write_traceability_matrix(matrix)

        return matrix

    def _extract_requirements(self, source_data: Dict) -> Dict[str, str]:
        """Extract requirements from source data"""
        requirements = {}

        # Look for documents with requirements
        for doc in source_data.get("documents", []):
            # If it's a Word document, look for requirement-like sections
            if doc.get("file_type") == ".docx":
                for section in doc.get("sections", []):
                    if any(req_word in section.get("heading", "").lower()
                           for req_word in ["requirement", "specification", "functional"]):
                        # Extract numbered requirements from the section content
                        section_content = section.get("content", "")
                        req_pattern = r'(REQ-\d+|R\d+|FR\d+|NFR\d+|[Rr]equirement\s+\d+)[\.:\)\s]+(.*?)(?=\n\n|\n[A-Z]|$)'

                        for match in re.finditer(req_pattern, section_content, re.DOTALL):
                            req_id = match.group(1).strip()
                            req_text = match.group(2).strip()
                            requirements[req_id] = req_text

        # If no structured requirements found, create some from the design docs
        if not requirements:
            req_id = 1
            for doc in source_data.get("documents", []):
                doc_content = doc.get("content", "")
                sentences = nltk.sent_tokenize(doc_content) if LIBRARIES_LOADED else doc_content.split('.')

                for sentence in sentences:
                    if any(req_word in sentence.lower()
                           for req_word in ["must", "should", "shall", "required", "needs to"]):
                        req_id_str = f"REQ-{req_id:03d}"
                        requirements[req_id_str] = sentence.strip()
                        req_id += 1

        return requirements

    def _extract_requirement_keywords(self, req_text: str) -> List[str]:
        """Extract key words from a requirement text"""
        if not LIBRARIES_LOADED:
            # Basic extraction - split and take non-stopwords
            words = req_text.split()
            # Simple English stopwords
            stopwords = ["the", "a", "an", "of", "in", "on", "at", "to", "for", "with", "by", "as", "is", "are"]
            return [word for word in words if word.lower() not in stopwords and len(word) > 3]

        # Use spaCy for better keyword extraction
        doc = nlp(req_text)
        keywords = []

        for token in doc:
            # Include nouns, verbs, adjectives, and proper nouns
            if token.pos_ in ["NOUN", "VERB", "ADJ", "PROPN"] and not token.is_stop:
                keywords.append(token.text)

        return keywords

    def _write_traceability_matrix(self, matrix: Dict):
        """Write traceability matrix to a file"""
        output_path = self.output_dir / "traceability_matrix.json"

        try:
            with open(output_path, 'w', encoding='utf-8') as f:
                json.dump(matrix, f, indent=2)

            logging.info(f"Wrote traceability matrix to {output_path}")

            # Also generate HTML report
            self._generate_html_report(matrix)

        except Exception as e:
            logging.error(f"Error writing traceability matrix: {e}")

    def _generate_html_report(self, matrix: Dict):
        """Generate HTML report for traceability matrix"""
        output_path = self.output_dir / "traceability_report.html"

        try:
            with open(output_path, 'w', encoding='utf-8') as f:
                f.write("""
                <!DOCTYPE html>
                <html>
                <head>
                    <title>Test Traceability Report</title>
                    <style>
                        body { font-family: Arial, sans-serif; margin: 20px; }
                        table { border-collapse: collapse; width: 100%; }
                        th, td { border: 1px solid #ddd; padding: 8px; text-align: left; }
                        th { background-color: #f2f2f2; }
                        tr:nth-child(even) { background-color: #f9f9f9; }
                        .covered { background-color: #d4edda; }
                        .not-covered { background-color: #f8d7da; }
                    </style>
                </head>
                <body>
                    <h1>Test Traceability Report</h1>
                """)

                # Requirements section
                f.write("<h2>Requirements</h2>")
                f.write("<table><tr><th>ID</th><th>Description</th></tr>")
                for req_id, req_text in matrix["requirements"].items():
                    f.write(f"<tr><td>{req_id}</td><td>{req_text}</td></tr>")
                f.write("</table>")

                # Test scenarios section
                f.write("<h2>Test Scenarios</h2>")
                f.write("<table><tr><th>Test Scenario</th><th>Requirements Covered</th></tr>")
                for scenario_name, covered_reqs in matrix["coverage"].items():
                    req_list = ", ".join(covered_reqs) if covered_reqs else "None"
                    coverage_class = "covered" if covered_reqs else "not-covered"
                    f.write(f"<tr class='{coverage_class}'><td>{scenario_name}</td><td>{req_list}</td></tr>")
                f.write("</table>")

                # Coverage matrix
                f.write("<h2>Coverage Matrix</h2>")
                f.write("<table><tr><th>Requirement ID</th>")
                for scenario in matrix["test_scenarios"]:
                    f.write(f"<th>{scenario}</th>")
                f.write("</tr>")

                for req_id in matrix["requirements"]:
                    f.write(f"<tr><td>{req_id}</td>")
                    for scenario in matrix["test_scenarios"]:
                        is_covered = req_id in matrix["coverage"].get(scenario, [])
                        coverage_mark = "✓" if is_covered else ""
                        cell_class = "covered" if is_covered else ""
                        f.write(f"<td class='{cell_class}'>{coverage_mark}</td>")
                    f.write("</tr>")
                f.write("</table>")

                f.write("</body></html>")

            logging.info(f"Wrote HTML traceability report to {output_path}")

        except Exception as e:
            logging.error(f"Error writing HTML report: {e}")


class LLMTestCaseGenerator:
    """Class to generate test cases using Azure OpenAI"""

    def __init__(self, api_key, model="gpt-4o", temperature=0.2):
        self.api_key = api_key
        self.model = model
        self.temperature = temperature
        # Azure endpoint configuration
        self.api_url = "https://llm-test-automation.openai.azure.com/openai/deployments/gpt-4o/chat/completions?api-version=2024-08-01-preview"
        self.headers = {"api-key": api_key, "Content-Type": "application/json"}
        # Track stats for reporting
        self.generated_scenarios_count = 0
        self.generated_steps_count = 0

    def call_azure_openai(self, messages, max_retries=5, base_wait=2):
        """Call Azure OpenAI API with retry logic"""
        import random
        import time
        from requests.exceptions import HTTPError

        for attempt in range(max_retries):
            try:
                response = requests.post(
                    self.api_url,
                    headers=self.headers,
                    json={
                        "messages": messages,
                        "temperature": self.temperature,
                        "top_p": 1.0,
                        "frequency_penalty": 0,
                        "presence_penalty": 0,
                        "max_tokens": 4000
                    }
                )

                # Check for rate limit before raising an exception
                if response.status_code == 429:
                    wait_time = base_wait * (2 ** attempt) + (random.uniform(0, 1))  # Add jitter
                    logging.warning(
                        f"Rate limit hit (429). Retrying in {wait_time:.2f} seconds... (Attempt {attempt + 1}/{max_retries})")
                    time.sleep(wait_time)
                    continue

                response.raise_for_status()
                return response.json()["choices"][0]["message"]["content"]

            except HTTPError as http_err:
                if response.status_code == 429:
                    wait_time = base_wait * (2 ** attempt) + (random.uniform(0, 1))
                    logging.warning(
                        f"Rate limit hit (429). Retrying in {wait_time:.2f} seconds... (Attempt {attempt + 1}/{max_retries})")
                    time.sleep(wait_time)
                else:
                    logging.error(f"HTTP error: {http_err}")
                    raise
            except Exception as e:
                logging.error(f"Unexpected error in OpenAI call: {e}")
                raise

        raise Exception("Max retries exceeded. OpenAI API is still rate limiting.")

    def generate_test_cases(self, context_data):
        """Generate test cases using the LLM.
        This method matches the expected signature in the framework.

        Args:
            context_data (dict): Input data for generating test cases.

        Returns:
            list: Generated test cases.
        """
        try:
            # Extract data from context_data
            requirements = context_data.get("requirements", [])
            scenarios = context_data.get("scenarios", [])
            additional_context = context_data.get("additional_context", "")

            # Construct the prompt
            prompt = f"""
            Generate test cases in Gherkin format based on the following requirements and scenarios:

            Requirements:
            {' '.join(requirements)}

            Scenarios:
            {' '.join(scenarios)}

            Additional Context:
            {additional_context}

            For each test case, include:
            1. A descriptive name
            2. Gherkin steps (Given, When, Then, And)
            3. Make steps detailed and specific

            Format your response as follows:
            Scenario: [Test Case Name]
            Given [precondition]
            When [action]
            Then [expected result]
            And [additional verification if needed]
            """

            # Call the API
            messages = [
                {"role": "system", "content": "You are a QA automation expert specializing in test case generation."},
                {"role": "user", "content": prompt}
            ]

            response_content = self.call_azure_openai(messages)

            # Process the response to extract test cases
            test_cases = self.extract_test_cases(response_content)

            return test_cases if test_cases else [self._create_fallback_test_case()]

        except Exception as e:
            logging.error(f"Error generating test cases with Azure OpenAI: {e}")
            return [self._create_fallback_test_case()]

    def extract_test_cases(self, response_content):
        """Extract test cases from the API response"""
        test_cases = []
        current_case = None

        for line in response_content.split('\n'):
            line = line.strip()
            if not line:
                continue

            # Check if this is a new scenario
            if line.startswith("Scenario:"):
                if current_case:
                    test_cases.append(current_case)
                current_case = {"name": line[9:].strip(), "steps": [], "description": ""}
            elif current_case and any(
                    line.startswith(step_type) for step_type in ["Given", "When", "Then", "And", "But"]):
                # This is a step
                for step_type in ["Given", "When", "Then", "And", "But"]:
                    if line.startswith(step_type):
                        description = line[len(step_type):].strip()
                        current_case["steps"].append({
                            "type": step_type,
                            "description": description
                        })
                        break
            elif current_case:
                # Additional info for the current case
                if current_case["description"]:
                    current_case["description"] += "\n" + line
                else:
                    current_case["description"] = line

        # Add the last test case
        if current_case:
            test_cases.append(current_case)

        return test_cases

    def _create_fallback_test_case(self):
        """Create a fallback test case when generation fails"""
        return {
            "name": "Basic Test Case",
            "description": "This is a fallback test case created when LLM generation failed.",
            "steps": [
                {"type": "Given", "description": "the system is ready"},
                {"type": "When", "description": "the user performs an action"},
                {"type": "Then", "description": "the system responds correctly"}
            ]
        }

    def generate_test_cases_only(self, context_data):
        """Generate test cases without detailed steps (first stage of two-stage process)"""
        try:
            # Extract data from context_data
            requirements = context_data.get("requirements", [])
            scenarios = context_data.get("scenarios", [])
            additional_context = context_data.get("additional_context", "")

            # Customize the prompt for test cases only (without detailed steps)
            prompt = f"""
            Generate test cases in Gherkin format based on the following requirements and scenarios.
            
            Focus ONLY on creating meaningful test case titles and descriptions. DO NOT create detailed steps.
            Use only placeholder steps (Given a setup, When an action occurs, Then a result is verified).
            
            The test steps will be created in a separate process, so focus on covering all test scenarios
            and edge cases implied by the requirements.
            
            Each test case should:
            1. Have a clear, descriptive name that explains what is being tested
            2. Include appropriate tags to categorize the test
            3. Include minimal placeholder steps (Given/When/Then) without details
            4. Cover positive scenarios, negative scenarios, and edge cases
            
            Requirements:
            {' '.join(requirements)}

            Scenarios:
            {' '.join(scenarios)}

            Additional Context:
            {additional_context}
            
            Format each test case as:
            
            @tag1 @tag2
            Scenario: Descriptive test case name
              Given a basic setup
              When the primary action occurs
              Then the expected result is verified
            """

            # Call the API
            messages = [
                {"role": "system", "content": "You are a QA automation expert specializing in test case generation."},
                {"role": "user", "content": prompt}
            ]

            response_content = self.call_azure_openai(messages)
            test_cases = self.extract_test_cases(response_content)
            
            # Add tags for better organization if not present
            for case in test_cases:
                if not case.get("tags"):
                    case["tags"] = ["automated"]
                
                # Ensure steps have minimal placeholder steps
                if not case.get("steps") or len(case.get("steps", [])) < 3:
                    case["steps"] = [
                        {"type": "Given", "description": "a basic setup"},
                        {"type": "When", "description": "the action is performed"},
                        {"type": "Then", "description": "the expected result is verified"}
                    ]
            
            self.generated_scenarios_count = len(test_cases)
            return test_cases if test_cases else [self._create_fallback_test_case()]

        except Exception as e:
            logging.error(f"Error generating test cases with Azure OpenAI: {e}")
            return [self._create_fallback_test_case()]

    def enhance_test_steps(self, test_cases, reference_data):
        """Enhance existing test cases with detailed steps (second stage of two-stage process)"""
        try:
            enhanced_test_cases = []
            total_steps = 0
            
            # Gather examples from reference data if available
            example_steps = reference_data.get("example_steps", [])
            design_context = reference_data.get("design_context", "")
            data_examples = reference_data.get("data_examples", [])
            
            # Create example steps text for the prompt
            example_steps_text = "\n".join(example_steps[:10]) if example_steps else ""
            
            # Process each test case
            for test_case in test_cases:
                # Format the original test case for the prompt
                scenario_name = test_case["name"]
                steps_text = ""
                for step in test_case.get("steps", []):
                    steps_text += f"{step['type']} {step['description']}\n"
                
                # Customize prompt for step enhancement
                prompt = f"""
                Enhance the following test case with detailed, specific test steps.
                
                The original test case has generic placeholder steps. Your task is to replace these
                with specific, actionable steps that a tester could follow precisely.
                
                TEST CASE:
                Scenario: {scenario_name}
                {steps_text}
                
                REFERENCE STEP EXAMPLES:
                {example_steps_text}
                
                DESIGN CONTEXT:
                {design_context[:2000]}
                
                DATA EXAMPLES:
                {', '.join(data_examples[:3])}
                
                Please provide detailed steps for this test case, following the same Gherkin format but with specific actions and verifications.
                Format your response as follows:
                Given [specific precondition with exact values or objects]
                When [specific action with exact parameters]
                Then [specific verification with expected results]
                """
                
                # Call the API
                messages = [
                    {"role": "system", "content": "You are a QA automation expert specializing in test step creation."},
                    {"role": "user", "content": prompt}
                ]
                
                response_content = self.call_azure_openai(messages)
                
                # Extract the enhanced steps
                enhanced_steps = []
                for line in response_content.strip().split('\n'):
                    line = line.strip()
                    if not line:
                        continue
                    
                    # Check if this is a step line
                    for step_type in ["Given", "When", "Then", "And", "But"]:
                        if line.startswith(step_type):
                            description = line[len(step_type):].strip()
                            
                            # Check if there's a data reference in JSON format
                            data_reference = None
                            json_match = re.search(r'(\{.*\})', description)
                            if json_match:
                                try:
                                    # Try to parse as JSON
                                    json_str = json_match.group(1)
                                    json.loads(json_str)  # Just to validate
                                    # Remove from description and store separately
                                    description = description.replace(json_str, "").strip()
                                    data_reference = json_str
                                except json.JSONDecodeError:
                                    # Not valid JSON, leave in description
                                    pass
                            
                            enhanced_steps.append({
                                "type": step_type,
                                "description": description,
                                "data_reference": data_reference
                            })
                            break
                
                # If no steps could be extracted, use the original placeholder steps
                if not enhanced_steps and test_case.get("steps"):
                    enhanced_steps = test_case["steps"]
                
                # Create the enhanced test case
                enhanced_test_case = test_case.copy()
                enhanced_test_case["steps"] = enhanced_steps
                
                # Track the number of steps
                total_steps += len(enhanced_steps)
                
                enhanced_test_cases.append(enhanced_test_case)
            
            self.generated_steps_count = total_steps
            return enhanced_test_cases
            
        except Exception as e:
            logging.error(f"Error enhancing test steps with Azure OpenAI: {e}")
            return test_cases  # Return the original test cases on error
            
    def generate_two_stage_test_cases(self, context_data, reference_data=None):
        """Generate test cases using a two-stage process for better quality"""
        # Stage 1: Generate test cases with placeholder steps
        test_cases = self.generate_test_cases_only(context_data)
        
        # If reference data is not provided, use empty defaults
        if not reference_data:
            reference_data = {
                "example_steps": [],
                "design_context": "",
                "data_examples": []
            }
        
        # Stage 2: Enhance test cases with detailed steps
        enhanced_test_cases = self.enhance_test_steps(test_cases, reference_data)
        
        return enhanced_test_cases
    
    def generate_report(self):
        """Generate a report of the test generation statistics"""
        return {
            "generated_scenarios": self.generated_scenarios_count,
            "generated_steps": self.generated_steps_count,
            "average_steps_per_scenario": self.generated_steps_count / self.generated_scenarios_count if self.generated_scenarios_count > 0 else 0
        }